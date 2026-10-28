"""Generate experiment report .docx files from requirements."""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_report(
    title: str,
    purpose: str,
    content: str,
    steps: list[str],
    requirements: list[str],
    output_path: Path,
    student_name: str = "",
    student_id: str = "",
) -> Path:
    """Generate a standard experiment report .docx.

    Args:
        title: Experiment title, e.g. "实验一 进程同步"
        purpose: Experiment purpose text
        content: Problem description / experiment content
        steps: List of experiment steps
        requirements: List of report requirements
        output_path: Where to save the .docx
        student_name: Student name (optional)
        student_id: Student ID (optional)

    Returns:
        Path to the generated file.
    """
    doc = Document()

    # set default font
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)

    # title
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # student info
    if student_name or student_id:
        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if student_name:
            info.add_run(f"姓名：{student_name}    ")
        if student_id:
            info.add_run(f"学号：{student_id}")

    # purpose
    doc.add_heading("实验目的", level=2)
    doc.add_paragraph(purpose)

    # content
    doc.add_heading("实验内容", level=2)
    doc.add_paragraph(content)

    # steps
    doc.add_heading("实验步骤", level=2)
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {step}")

    # implementation
    doc.add_heading("实验实现", level=2)
    doc.add_paragraph("（此处填写算法伪代码与实现代码）")
    doc.add_paragraph("")
    doc.add_paragraph("算法伪代码：")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("实现代码：")
    doc.add_paragraph("")
    doc.add_paragraph("")

    # analysis
    doc.add_heading("结果分析", level=2)
    doc.add_paragraph("（此处填写时间复杂度分析、测试用例结果）")
    doc.add_paragraph("")
    doc.add_paragraph("时间复杂度分析：")
    doc.add_paragraph("")
    doc.add_paragraph("测试用例结果：")
    doc.add_paragraph("")

    # summary
    doc.add_heading("实验总结", level=2)
    doc.add_paragraph("（此处填写实验中的关键问题和难点总结）")

    # requirements (reference)
    doc.add_heading("实验报告要求", level=2)
    for req in requirements:
        doc.add_paragraph(f"• {req}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


# Standard template for algorithm course experiments
ALGORITHM_STEPS = [
    "根据实验内容设计算法伪代码进行算法描述",
    "利用C++/C/Java等编程语言对算法伪代码进行工程化实现",
    "输入测试用例对算法进行验证",
    "分析算法时间复杂度",
]

ALGORITHM_REQUIREMENTS = [
    "结合题目要求给出格式正确、注释清晰的算法伪代码",
    "将实现正确、运行无误的实现代码以源文件形式提交",
    "给出算法时间复杂度的完整分析和推导过程",
    "总结实验中的关键问题和难点",
]


def generate_algorithm_report(
    title: str,
    purpose: str,
    content: str,
    output_path: Path,
    student_name: str = "",
    student_id: str = "",
) -> Path:
    """Convenience wrapper for algorithm course experiments."""
    return generate_report(
        title=title,
        purpose=purpose,
        content=content,
        steps=ALGORITHM_STEPS,
        requirements=ALGORITHM_REQUIREMENTS,
        output_path=output_path,
        student_name=student_name,
        student_id=student_id,
    )
